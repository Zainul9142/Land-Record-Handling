import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_sih_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Official Colors
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_DARK_TEXT = RGBColor(15, 23, 42)       # Charcoal/Slate 900
    COLOR_NAVY = RGBColor(15, 61, 107)           # SIH Header Navy
    COLOR_BLUE_SUB = RGBColor(0, 86, 150)        # Vibrant Blue (❖ Proposed Solution)
    COLOR_FOOTER_BLUE = RGBColor(30, 96, 145)    # Official Template Footer Blue
    COLOR_CARD_BG = RGBColor(248, 250, 252)      # Slate 50
    COLOR_CARD_BORDER = RGBColor(203, 213, 225)  # Slate 300
    COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)    # Blue 600
    COLOR_AMBER = RGBColor(217, 119, 6)          # Amber 600
    COLOR_EMERALD = RGBColor(16, 185, 129)       # Emerald 500
    COLOR_MUTED = RGBColor(71, 85, 105)          # Slate 600

    logo_path = r"d:\SIH\1\backend\extracted_assets\sih_logo_clean.png"
    bulb_path = r"d:\SIH\1\backend\extracted_assets\brain_bulb_clean_final.png"

    has_logo = os.path.exists(logo_path)
    has_bulb = os.path.exists(bulb_path)

    def set_white_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_WHITE

    def add_common_template_elements(slide, slide_num, title_text):
        # 1. Top Left Oval Badge: "TEAM AGNI"
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.25), Inches(1.8), Inches(1.05))
        oval.fill.solid()
        oval.fill.fore_color.rgb = COLOR_WHITE
        oval.line.color.rgb = COLOR_DARK_TEXT
        oval.line.width = Pt(1.5)
        tf_o = oval.text_frame
        tf_o.word_wrap = True
        p_o1 = tf_o.paragraphs[0]
        p_o1.text = "TEAM"
        p_o1.font.bold = True
        p_o1.font.size = Pt(12)
        p_o1.font.color.rgb = COLOR_NAVY
        p_o1.alignment = PP_ALIGN.CENTER
        
        p_o2 = tf_o.add_paragraph()
        p_o2.text = "AGNI"
        p_o2.font.bold = True
        p_o2.font.size = Pt(13)
        p_o2.font.color.rgb = COLOR_AMBER
        p_o2.alignment = PP_ALIGN.CENTER

        # 2. Top Center Slide Title
        t_box = slide.shapes.add_textbox(Inches(2.3), Inches(0.25), Inches(8.7), Inches(0.95))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.bold = True
        p_t.font.size = Pt(21)
        p_t.font.color.rgb = COLOR_DARK_TEXT
        p_t.alignment = PP_ALIGN.CENTER

        # 3. Top Right SIH Logo
        if has_logo:
            slide.shapes.add_picture(logo_path, Inches(11.2), Inches(0.2), width=Inches(1.75))

        # 4. Bottom Footer Strip across width
        footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5))
        footer.fill.solid()
        footer.fill.fore_color.rgb = COLOR_FOOTER_BLUE
        footer.line.fill.background()
        
        tf_f = footer.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = f"@SIH Idea submission- Template {slide_num}"
        p_f.font.size = Pt(12)
        p_f.font.color.rgb = COLOR_WHITE
        p_f.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 1: TITLE PAGE (Exact Match to SIH Template)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide1)

    # Top SIH Header
    header_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(9.5), Inches(0.8))
    p_h = header_box.text_frame.paragraphs[0]
    p_h.text = "SMART INDIA HACKATHON 2026"
    p_h.font.bold = True
    p_h.font.size = Pt(32)
    p_h.font.color.rgb = COLOR_NAVY

    # Top Right SIH Logo
    if has_logo:
        slide1.shapes.add_picture(logo_path, Inches(11.0), Inches(0.2), width=Inches(1.9))

    # Centered: TITLE PAGE
    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.5))
    p_sub = sub_box.text_frame.paragraphs[0]
    p_sub.text = "TITLE PAGE"
    p_sub.font.bold = True
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = COLOR_DARK_TEXT
    p_sub.alignment = PP_ALIGN.CENTER

    # Main Left Box with Exact Template Pointers
    left_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.6), Inches(5.2))
    tf1 = left_box.text_frame
    tf1.word_wrap = True

    fields = [
        ("• Problem Statement ID – ", "SIH26014", COLOR_AMBER),
        ("• Problem Statement Title- ", "BhoomiShield: Pan-India AI Land Record Governance, Cadastral GIS Mapping & Multi-Jurisdictional Fraud Prevention Engine", COLOR_NAVY),
        ("• Theme- ", "Smart Governance / Digital India / Land Administration Modernization", COLOR_DARK_TEXT),
        ("• PS Category- Software/Hardware ", "Software", COLOR_ACCENT_BLUE),
        ("• Team ID- ", "2404921540161", COLOR_DARK_TEXT),
        ("• Team Name (Registered on portal) ", "TEAM AGNI (Leader: Priyanshu kumar yadav)", COLOR_AMBER)
    ]

    for idx, (lbl, val, col) in enumerate(fields):
        p = tf1.paragraphs[0] if idx == 0 else tf1.add_paragraph()
        run1 = p.add_run()
        run1.text = lbl
        run1.font.bold = True
        run1.font.size = Pt(13.5)
        run1.font.color.rgb = COLOR_DARK_TEXT

        run2 = p.add_run()
        run2.text = val
        run2.font.bold = True
        run2.font.size = Pt(13.5)
        run2.font.color.rgb = col
        p.space_after = Pt(12)

    # Right Side Graphic: Clean Brain Bulb Artwork
    if has_bulb:
        slide1.shapes.add_picture(bulb_path, Inches(8.8), Inches(1.7), width=Inches(3.6))

    # ==========================================
    # SLIDE 2: IDEA TITLE & PROPOSED SOLUTION
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide2)
    add_common_template_elements(slide2, 2, "BHOOMISHIELD: PAN-INDIA LAND RECORD AI & GOVERNANCE PLATFORM")

    sub2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.4))
    p_s2 = sub2.text_frame.paragraphs[0]
    p_s2.text = "❖ Proposed Solution (Describe your Idea/Solution/Prototype)"
    p_s2.font.bold = True
    p_s2.font.size = Pt(15)
    p_s2.font.color.rgb = COLOR_BLUE_SUB

    s2_cards = [
        ("• Detailed Explanation of Proposed Solution", COLOR_NAVY, [
            "Unified Pan-India Aggregation: Standardizes fragmented state land data (Khatian, RoR, Register-II, Jamabandi, Deeds, Mutation logs) into an interoperable ULPIN canonical schema.",
            "Interactive Cadastral GIS & Location Search: Instant 1-click parcel lookup by GPS coordinates, satellite boundary polygons, Khata/Plot number, or owner name.",
            "Bhoomi Vault & Authority Portal: Secure citizen repository for land deeds and tax receipts paired with real-time revenue officer mutation & dispute dashboards."
        ]),
        ("• How It Addresses the Problem", COLOR_AMBER, [
            "Eliminates Portal Fragmentation: Replaces 30+ disparate regional portals with a centralized, unified intelligent query and verification interface.",
            "Prevents Fraudulent Resales: Instantly flags unmutated deeds, pending bank mortgages, and active civil court stay orders before financial transaction.",
            "Protects Tribal & Protected Land: Enforces statutory tenancy regulations (e.g. CNT/SPT Act, UP Revenue Code §98, Karnataka PTCL) preventing illegal land alienation."
        ]),
        ("• Innovation and Uniqueness of Solution", COLOR_EMERALD, [
            "Deterministic 0–100 AI Risk Engine: Heuristic rule-based scoring (R001–R007) ensuring transparent, fully explainable risk factors with zero hallucination.",
            "Multilingual Voice Legal Assistant: Conversational legal aid and dispute drafting in 11 Indian regional languages for high rural accessibility.",
            "Cryptographic SHA-256 Audit Trail: Generates QR-coded official Land Health Certificates verifying textual & spatial consistency across state archives."
        ])
    ]

    for idx, (header_text, header_color, points) in enumerate(s2_cards):
        left_pos = Inches(0.8 + idx * 3.95)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), Inches(3.8), Inches(5.05))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = header_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)

        p_h = tf.paragraphs[0]
        p_h.text = header_text
        p_h.font.bold = True
        p_h.font.size = Pt(11.5)
        p_h.font.color.rgb = header_color
        p_h.space_after = Pt(8)

        for pt in points:
            p_pt = tf.add_paragraph()
            parts = pt.split(":", 1)
            if len(parts) == 2:
                r1 = p_pt.add_run()
                r1.text = "✔ " + parts[0] + ":"
                r1.font.bold = True
                r1.font.size = Pt(10)
                r1.font.color.rgb = COLOR_DARK_TEXT
                
                r2 = p_pt.add_run()
                r2.text = parts[1]
                r2.font.size = Pt(9.5)
                r2.font.color.rgb = COLOR_MUTED
            else:
                p_pt.text = "✔ " + pt
                p_pt.font.size = Pt(9.5)
                p_pt.font.color.rgb = COLOR_MUTED
            p_pt.space_after = Pt(8)

    # ==========================================
    # SLIDE 3: TECHNICAL APPROACH
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide3)
    add_common_template_elements(slide3, 3, "TECHNICAL APPROACH")

    left_s3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.4))
    left_s3.fill.solid()
    left_s3.fill.fore_color.rgb = COLOR_CARD_BG
    left_s3.line.color.rgb = COLOR_NAVY
    left_s3.line.width = Pt(1.5)

    tf_s3_l = left_s3.text_frame
    tf_s3_l.word_wrap = True
    tf_s3_l.margin_left = Inches(0.25)
    tf_s3_l.margin_right = Inches(0.25)
    tf_s3_l.margin_top = Inches(0.2)

    p_tech_hdr = tf_s3_l.paragraphs[0]
    p_tech_hdr.text = "• Technologies to be used (e.g. languages, frameworks, hardware)"
    p_tech_hdr.font.bold = True
    p_tech_hdr.font.size = Pt(12)
    p_tech_hdr.font.color.rgb = COLOR_NAVY
    p_tech_hdr.space_after = Pt(6)

    tech_stack = [
        ("Frontend Application", "React 18, TypeScript, Vite, Tailwind CSS, Lucide Icons, Mobile Responsive PWA architecture."),
        ("Backend & Microservices", "FastAPI (Python 3.11+), Asynchronous REST Endpoints, Uvicorn, Pydantic v2 validation."),
        ("Cadastral GIS & Spatial", "Leaflet.js, Esri Satellite Imagery, OpenStreetMap, Point-in-Polygon (Ray Casting) & Turf.js / Shapely."),
        ("Database & Hashing", "SQLite / PostgreSQL with Spatial Indexing, Document Store, Cryptographic SHA-256 integrity checks."),
        ("AI & Multilingual NLP", "Custom 7-Rule Deterministic Risk Engine, Speech Recognition API & Web Speech Synthesis (11 Indian Languages)."),
        ("Report Generation", "ReportLab Automated PDF Engine with dynamic QR Code verification & bilingual summaries.")
    ]

    for cat, desc in tech_stack:
        p_t = tf_s3_l.add_paragraph()
        r1 = p_t.add_run()
        r1.text = "▪ " + cat + ": "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_ACCENT_BLUE

        r2 = p_t.add_run()
        r2.text = desc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_DARK_TEXT
        p_t.space_after = Pt(5)

    right_s3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.45), Inches(5.9), Inches(5.4))
    right_s3.fill.solid()
    right_s3.fill.fore_color.rgb = RGBColor(241, 245, 249)
    right_s3.line.color.rgb = COLOR_AMBER
    right_s3.line.width = Pt(1.5)

    tf_s3_r = right_s3.text_frame
    tf_s3_r.word_wrap = True
    tf_s3_r.margin_left = Inches(0.25)
    tf_s3_r.margin_right = Inches(0.25)
    tf_s3_r.margin_top = Inches(0.2)

    p_meth_hdr = tf_s3_r.paragraphs[0]
    p_meth_hdr.text = "• Methodology & Process for Implementation (Working Prototype)"
    p_meth_hdr.font.bold = True
    p_meth_hdr.font.size = Pt(12)
    p_meth_hdr.font.color.rgb = COLOR_AMBER
    p_meth_hdr.space_after = Pt(6)

    method_steps = [
        ("Step 1: Multi-State Ingestion Pipeline", "State land portals (e.g. Jharbhoomi, Bhulekh UP, Mahabhumi, BanglarBhumi) ingested into canonical schema matching ULPIN standard."),
        ("Step 2: Cadastral Polygon & Location Binding", "Associates survey numbers with GIS polygon geometries. Live GPS / click queries locate parcel boundaries on satellite map."),
        ("Step 3: Multi-Layer Deterministic Risk Engine", "Evaluates R001 (Owner Mismatch), R002 (Area Variance), R003 (Missing Geo), R004/R005 (Mutation SLA), R006 (Deed Mismatch), R007 (Court Litigation)."),
        ("Step 4: Bhoomi Vault & Citizen Governance", "Citizens securely upload land records; officials triage grievances; real-time tamper-proof PDF audit certificates generated instantly.")
    ]

    for stitle, sdesc in method_steps:
        p_m = tf_s3_r.add_paragraph()
        r1 = p_m.add_run()
        r1.text = "▶ " + stitle + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_NAVY

        r2 = p_m.add_run()
        r2.text = sdesc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_DARK_TEXT
        p_m.space_after = Pt(6)

    # ==========================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide4)
    add_common_template_elements(slide4, 4, "FEASIBILITY AND VIABILITY")

    s4_sections = [
        ("• Analysis of the Feasibility of the Idea", COLOR_NAVY, [
            "Technical Feasibility: Fully functional working prototype deployed with 10,000+ parcels across Jharkhand, UP, Maharashtra, and pan-India states with sub-150ms query latency.",
            "Operational Feasibility: Non-invasive architecture that operates as an analytical intelligence layer over existing state databases without requiring costly database re-engineering.",
            "Financial Viability: Built on high-performance open-source stack (FastAPI, React, SQLite/PostgreSQL, Leaflet) minimizing infrastructure and licensing overhead."
        ]),
        ("• Potential Challenges and Risks", COLOR_AMBER, [
            "Heterogeneous Terminology: Varied state revenue terminologies (Khatian, RoR, 7/12, Jamabandi, Patta, Chitta, Dag) across different state jurisdictions.",
            "State Portal Downtime & Anti-Scraping: Legacy government servers experience intermittent downtime and strict IP rate-limiting during peak daytime traffic.",
            "Digital Literacy in Rural India: Complex land laws and technical legal terminology pose adoption hurdles for rural farmers and non-English speakers."
        ]),
        ("• Strategies for Overcoming These Challenges", COLOR_EMERALD, [
            "Universal Canonical Data Adapter: Standardizes diverse state terminologies into a unified Bhu-Aadhaar / ULPIN compliant schema automatically.",
            "Asynchronous Caching & Resilient Adapters: Multi-tiered local SQLite/Redis caching layer with graceful fallback simulation to guarantee 99.9% uptime.",
            "Multilingual Voice AI & Color Indicators: Interactive voice assistance in 11 vernacular Indian languages and intuitive Green/Yellow/Red risk badges for instant clarity."
        ])
    ]

    for idx, (hdr_text, hdr_color, pts) in enumerate(s4_sections):
        left_p = Inches(0.8 + idx * 3.95)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_p, Inches(1.45), Inches(3.8), Inches(5.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = hdr_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)

        p_h = tf.paragraphs[0]
        p_h.text = hdr_text
        p_h.font.bold = True
        p_h.font.size = Pt(11.5)
        p_h.font.color.rgb = hdr_color
        p_h.space_after = Pt(8)

        for pt in points:
            p_pt = tf.add_paragraph()
            parts = pt.split(":", 1)
            if len(parts) == 2:
                r1 = p_pt.add_run()
                r1.text = "✔ " + parts[0] + ":"
                r1.font.bold = True
                r1.font.size = Pt(10)
                r1.font.color.rgb = COLOR_DARK_TEXT
                
                r2 = p_pt.add_run()
                r2.text = parts[1]
                r2.font.size = Pt(9.5)
                r2.font.color.rgb = COLOR_MUTED
            else:
                p_pt.text = "✔ " + pt
                p_pt.font.size = Pt(9.5)
                p_pt.font.color.rgb = COLOR_MUTED
            p_pt.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: IMPACT AND BENEFITS
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide5)
    add_common_template_elements(slide5, 5, "IMPACT AND BENEFITS")

    left_s5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.4))
    left_s5.fill.solid()
    left_s5.fill.fore_color.rgb = COLOR_CARD_BG
    left_s5.line.color.rgb = COLOR_NAVY
    left_s5.line.width = Pt(1.5)

    tf_s5_l = left_s5.text_frame
    tf_s5_l.word_wrap = True
    tf_s5_l.margin_left = Inches(0.25)
    tf_s5_l.margin_right = Inches(0.25)
    tf_s5_l.margin_top = Inches(0.2)

    p_imp_hdr = tf_s5_l.paragraphs[0]
    p_imp_hdr.text = "• Potential Impact on the Target Audience"
    p_imp_hdr.font.bold = True
    p_imp_hdr.font.size = Pt(13)
    p_imp_hdr.font.color.rgb = COLOR_NAVY
    p_imp_hdr.space_after = Pt(6)

    audiences = [
        ("Citizens & Rural Farmers", "Prevents fraudulent land dispossession, unauthorized transfers, and opaque broker commissions with instant verified ownership history."),
        ("Property Buyers & Developers", "Conducts instantaneous 360° title due-diligence, mortgage checks, and litigation status before capital investment."),
        ("Banks & Financial Institutions", "Accelerates loan origination and mortgage collateral verification from weeks to seconds while reducing non-performing assets."),
        ("Revenue Officers & Administration", "Provides automated SLA tracking for pending mutation files, grievance analytics, and rapid dispute resolution workflows.")
    ]

    for aud, adesc in audiences:
        p_a = tf_s5_l.add_paragraph()
        r1 = p_a.add_run()
        r1.text = "👥 " + aud + ": "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_ACCENT_BLUE

        r2 = p_a.add_run()
        r2.text = adesc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_DARK_TEXT
        p_a.space_after = Pt(5)

    right_s5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.45), Inches(5.9), Inches(5.4))
    right_s5.fill.solid()
    right_s5.fill.fore_color.rgb = RGBColor(241, 245, 249)
    right_s5.line.color.rgb = COLOR_EMERALD
    right_s5.line.width = Pt(1.5)

    tf_s5_r = right_s5.text_frame
    tf_s5_r.word_wrap = True
    tf_s5_r.margin_left = Inches(0.25)
    tf_s5_r.margin_right = Inches(0.25)
    tf_s5_r.margin_top = Inches(0.2)

    p_ben_hdr = tf_s5_r.paragraphs[0]
    p_ben_hdr.text = "• Benefits of the Solution (Social, Economic, Environmental)"
    p_ben_hdr.font.bold = True
    p_ben_hdr.font.size = Pt(13)
    p_ben_hdr.font.color.rgb = COLOR_EMERALD
    p_ben_hdr.space_after = Pt(6)

    benefits = [
        ("Social Equity & Tenancy Protection", "Safeguards tribal and marginalized farmers under CNT/SPT Acts and UP Revenue Code §98; democratizes legal literacy across rural Bharat."),
        ("Economic Acceleration", "Unlocks locked economic capital in disputed land parcels; directly targets India's 66% civil court land litigation backlog."),
        ("Environmental & Wetland Protection", "Cadastral satellite overlay automatically identifies illegal encroachments on forest land, waterbodies (Jalasay), and flood plains."),
        ("Governance Transparency", "End-to-end cryptographic audit trails (SHA-256) eliminate document tampering and backdated mutation entries.")
    ]

    for bcat, bdesc in benefits:
        p_b = tf_s5_r.add_paragraph()
        r1 = p_b.add_run()
        r1.text = "🌟 " + bcat + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_NAVY

        r2 = p_b.add_run()
        r2.text = bdesc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_DARK_TEXT
        p_b.space_after = Pt(6)

    # ==========================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide6)
    add_common_template_elements(slide6, 6, "RESEARCH AND REFERENCES")

    card_s6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.4))
    card_s6.fill.solid()
    card_s6.fill.fore_color.rgb = COLOR_CARD_BG
    card_s6.line.color.rgb = COLOR_NAVY
    card_s6.line.width = Pt(1.5)

    tf_s6 = card_s6.text_frame
    tf_s6.word_wrap = True
    tf_s6.margin_left = Inches(0.3)
    tf_s6.margin_right = Inches(0.3)
    tf_s6.margin_top = Inches(0.2)

    p_ref_hdr = tf_s6.paragraphs[0]
    p_ref_hdr.text = "• Details / Links of the Reference and Research Work"
    p_ref_hdr.font.bold = True
    p_ref_hdr.font.size = Pt(13.5)
    p_ref_hdr.font.color.rgb = COLOR_NAVY
    p_ref_hdr.space_after = Pt(8)

    references = [
        ("1. Digital India Land Records Modernization Programme (DILRMP)", 
         "Department of Land Resources (DoLR), Ministry of Rural Development, Government of India.\nFramework for computerization of land records, cadastral map digitization, and survey integration."),
        ("2. Unique Land Parcel Identification Number (ULPIN) / Bhu-Aadhaar Guidelines (2021-2024)",
         "National standards for assigning 14-digit alphanumeric geo-referenced identification to each land parcel across India."),
        ("3. Statutory State Revenue Acts & Tenancy Codes Studied",
         "• Chota Nagpur Tenancy Act (CNT), 1908 & Santhal Parganas Tenancy Act (SPT), 1949 (Jharkhand)\n• Uttar Pradesh Revenue Code, 2006 (Section 98 - Tenancy Restrictions & Mutation SLAs)\n• Maharashtra Land Revenue Code (MLRC), 1966 & Karnataka PTCL Act, 1978\n• Real Estate (Regulation and Development) Act (RERA), 2016"),
        ("4. NITI Aayog Report on Land Governance in India & Judicial Reforms",
         "Research findings on title verification bottlenecks, land dispute litigation resolution, and automated conclusive titling systems."),
        ("5. Open-Source Implementation & Live Prototype Repository",
         "• GitHub Codebase: https://github.com/Zainul9142/Land-Record-Handling\n• Modules Deployed: Multi-State Ingestion Pipeline, Cadastral GIS Leaflet Map, 7-Rule Heuristic Engine, 11-Language Multilingual Localization, Citizen Bhoomi Vault & Official Officer Portal.")
    ]

    for title, body in references:
        p_r = tf_s6.add_paragraph()
        r1 = p_r.add_run()
        r1.text = title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_ACCENT_BLUE

        r2 = p_r.add_run()
        r2.text = body
        r2.font.size = Pt(9)
        r2.font.color.rgb = COLOR_DARK_TEXT
        p_r.space_after = Pt(6)

    output_path = r"d:\SIH\1\BhoomiShield_SIH2026_TEAM_AGNI.pptx"
    prs.save(output_path)
    print(f"Successfully generated SIH 2026 PPTX presentation at: {output_path}")

if __name__ == "__main__":
    build_sih_presentation()
